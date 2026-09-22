"""Normalize monthly mutual-fund portfolio workbooks into one tabular dataset."""

from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


OUTPUT_COLUMNS = [
    "fund_isin", "fund_name", "portfolio_date", "record_type", "section",
    "instrument_name", "instrument_isin", "isin_status", "industry_or_rating", "quantity",
    "market_value_lakh", "market_value_inr", "pct_of_nav", "yield_pct",
    "yield_to_call_pct", "coupon_pct", "position", "price_when_purchased",
    "current_price", "margin_maintained_lakh", "source_file", "source_sheet",
    "source_row", "raw_row_json",
]

ALIASES = {
    "instrument_name": {"companyissuerinstrumentname", "nameofinstrument", "nameoftheinstrument", "nameofinstrumentissuer", "nameoftheinstrumentissuer", "underlyingsecurityname"},
    "instrument_isin": {"isin", "isincode"},
    "industry_or_rating": {"industry", "industryrating", "ratingindustry"},
    "quantity": {"quantity"},
    "market_value_lakh": {"marketvalue", "marketvaluersinlakhs", "marketvaluersinlacs", "marketfairvaluersinlacs", "exposuremarketvaluerslakh"},
    "pct_of_nav": {"tonav", "tonetassets", "toaum"},
    "yield_pct": {"yield", "ytm"},
    "yield_to_call_pct": {"yieldtocall", "ytc"},
    "coupon_pct": {"coupon"},
    "position": {"longshort"},
    "price_when_purchased": {"futurepricewhenpurchasedinrs"},
    "current_price": {"currentpriceofthecontractinrs"},
    "margin_maintained_lakh": {"marginmaintainedrsinlakhs"},
}


def clean_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    return text or None


def clean_instrument_name(value: Any) -> str | None:
    text = clean_text(value)
    if not text:
        return None
    text = re.sub(r"\s+\$+\s*$", "", text)
    text = re.sub(r"^EQ\s*[-:]\s*", "", text, flags=re.IGNORECASE)
    return text.strip(" .") or None


def clean_isin(value: Any) -> str | None:
    text = clean_text(value)
    if not text:
        return None
    match = re.search(r"\b([A-Z]{2}[A-Z0-9]{9}[0-9])\b", text.upper().replace(" ", ""))
    return match.group(1) if match else None


def parse_number(value: Any) -> float | int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    text = clean_text(value)
    if not text or text.upper() in {"-", "—", "NIL", "NA", "N/A"}:
        return None
    try:
        number = float(text.replace(",", "").replace("%", ""))
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        return None
    for pattern in (r"\b\d{4}-\d{2}-\d{2}\b", r"\b\d{1,2}[-/]\d{1,2}[-/]\d{4}\b", r"\b[A-Za-z]{3,9}\s+\d{1,2},?\s*\d{4}\b"):
        match = re.search(pattern, text)
        if not match:
            continue
        candidate = " ".join(match.group(0).replace(",", " ").split())
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%B %d %Y", "%b %d %Y"):
            try:
                return datetime.strptime(candidate, fmt).date()
            except ValueError:
                pass
    return None


def percent_value(value: Any) -> float | int | None:
    number = parse_number(value)
    if number is None:
        return None
    return number * 100 if abs(number) <= 1 else number


def find_header(row: tuple[Any, ...]) -> dict[str, int] | None:
    mapped: dict[str, int] = {}
    for index, value in enumerate(row):
        key = clean_key(value)
        inferred_field = None
        if "market" in key and "value" in key and ("lakh" in key or "lac" in key):
            inferred_field = "market_value_lakh"
        elif "percent" in key or key in {"tonav", "tonetassets", "toaum"}:
            inferred_field = "pct_of_nav"
        for field, aliases in ALIASES.items():
            if (key in aliases or field == inferred_field) and field not in mapped:
                mapped[field] = index
    portfolio_fields = {"instrument_isin", "quantity", "market_value_lakh", "pct_of_nav"}
    derivative_fields = {"position", "price_when_purchased", "current_price", "margin_maintained_lakh"}
    return mapped if "instrument_name" in mapped and (len(set(mapped) & portfolio_fields) >= 2 or len(set(mapped) & derivative_fields) >= 2) else None


def workbook_metadata(path: Path, rows: Iterable[tuple[Any, ...]]) -> tuple[str, str, date | None]:
    fund_isin = path.name.split("_", 1)[0].upper()
    fund_name = clean_text(path.stem.split("_", 1)[1] if "_" in path.stem else path.stem) or path.stem
    report_date = None
    for row in rows:
        for index, value in enumerate(row):
            text = clean_text(value)
            if text and text.lower().startswith("scheme name") and index + 1 < len(row):
                fund_name = clean_text(row[index + 1]) or fund_name
            report_date = report_date or parse_date(value)
    return fund_isin, fund_name, report_date


def normalize_workbook(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    metadata_rows = []
    for sheet in workbook.worksheets:
        metadata_rows.extend(list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), values_only=True)))
    fund_isin, fund_name, portfolio_date = workbook_metadata(path, metadata_rows)
    records: list[dict[str, Any]] = []

    for sheet in workbook.worksheets:
        header_map = None
        section = None
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if header_map is None:
                candidate = find_header(row)
                if candidate is None:
                    continue
                header_map = candidate
                continue
            values = {field: row[index] if index < len(row) else None for field, index in header_map.items()}
            if not clean_text(values.get("instrument_name")) and "instrument_isin" in header_map:
                isin_index = header_map["instrument_isin"]
                shifted_names = [clean_text(value) for value in row[:isin_index] if clean_text(value)]
                if shifted_names:
                    values["instrument_name"] = shifted_names[-1]
            instrument = clean_instrument_name(values.get("instrument_name"))
            has_data = any(values.get(field) not in (None, "", "-") for field in header_map if field != "instrument_name")
            if instrument and not has_data:
                section = instrument
                continue
            instrument_isin = clean_isin(values.get("instrument_isin"))
            numeric_fields = ("quantity", "market_value_lakh", "price_when_purchased", "current_price", "margin_maintained_lakh")
            has_numeric_position = any(parse_number(values.get(field)) is not None for field in numeric_fields)
            if not instrument or not has_data:
                continue
            record_type = "derivative" if "position" in header_map else "portfolio"
            if record_type == "portfolio" and not instrument_isin and parse_number(values.get("quantity")) is None:
                continue
            if record_type == "derivative" and not has_numeric_position:
                continue
            market_value_lakh = parse_number(values.get("market_value_lakh"))
            raw_values = [clean_text(value) if not isinstance(value, (int, float, datetime, date)) else value for value in row]
            records.append({
                "fund_isin": fund_isin, "fund_name": fund_name,
                "portfolio_date": portfolio_date.isoformat() if portfolio_date else None,
                "record_type": record_type, "section": section,
                "instrument_name": instrument, "instrument_isin": instrument_isin,
                "isin_status": "present" if instrument_isin else "not_disclosed",
                "industry_or_rating": clean_text(values.get("industry_or_rating")),
                "quantity": parse_number(values.get("quantity")),
                "market_value_lakh": market_value_lakh,
                "market_value_inr": market_value_lakh * 100000 if market_value_lakh is not None else None,
                "pct_of_nav": percent_value(values.get("pct_of_nav")),
                "yield_pct": percent_value(values.get("yield_pct")),
                "yield_to_call_pct": percent_value(values.get("yield_to_call_pct")),
                "coupon_pct": percent_value(values.get("coupon_pct")),
                "position": clean_text(values.get("position")),
                "price_when_purchased": parse_number(values.get("price_when_purchased")),
                "current_price": parse_number(values.get("current_price")),
                "margin_maintained_lakh": parse_number(values.get("margin_maintained_lakh")),
                "source_file": path.name, "source_sheet": sheet.title, "source_row": row_number,
                "raw_row_json": json.dumps(raw_values, default=str, ensure_ascii=True),
            })
    workbook.close()
    return records


def write_outputs(records: list[dict[str, Any]], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "normalized_monthly_portfolios.csv"
    try:
        handle = csv_path.open("w", newline="", encoding="utf-8")
    except PermissionError:
        csv_path = output_dir / "normalized_monthly_portfolios_clean.csv"
        handle = csv_path.open("w", newline="", encoding="utf-8")
    with handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(records)
    try:
        import pandas as pd
        pd.DataFrame(records, columns=OUTPUT_COLUMNS).to_parquet(output_dir / "normalized_monthly_portfolios.parquet", index=False)
    except (ImportError, ValueError, OSError):
        pass


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=Path("data/raw/Funds Monthly Portfolio Disclosure(18 funds)/current"))
    parser.add_argument("--output-dir", type=Path, default=Path("data/processed"))
    args = parser.parse_args()
    paths = sorted(args.input_dir.glob("*.xlsx"))
    records = [record for path in paths for record in normalize_workbook(path)]
    write_outputs(records, args.output_dir)
    print(f"Normalized {len(records):,} records from {len(paths)} workbooks into {args.output_dir}")


if __name__ == "__main__":
    main()